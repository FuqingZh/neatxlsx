from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

import neatxlsx as nx
import openpyxl
import polars as pl
import pytest


def test_m1_value_kinds_precision_and_temporal_fallbacks(tmp_path: Path) -> None:
    output = tmp_path / "m1.xlsx"
    data = pl.DataFrame(
        {
            "flag": pl.Series([True, False], dtype=pl.Boolean),
            "identifier": pl.Series(
                [123456789012345, 1234567890123456], dtype=pl.Int64
            ),
            "decimal": pl.Series(
                [Decimal("1.20"), Decimal("123456789012345678.90")],
                dtype=pl.Decimal(20, 2),
            ),
            "date": pl.Series([date(2024, 1, 1), date(1899, 12, 31)], dtype=pl.Date),
            "datetime": pl.Series(
                [
                    datetime(2024, 1, 1, 1, 2, 3, 123000),
                    datetime(2024, 1, 1, 1, 2, 3, 123456),
                ],
                dtype=pl.Datetime("us"),
            ),
            "time": pl.Series(
                [time(1, 2, 3, 123000), time(1, 2, 3, 123456)], dtype=pl.Time
            ),
            "duration": pl.Series(
                [timedelta(seconds=1), timedelta(microseconds=-1000)],
                dtype=pl.Duration("us"),
            ),
            "literal": ["=1+1", "plain"],
        }
    )

    with nx.Workbook(output) as workbook:
        workbook.write_sheet(data, "Report", autofit=nx.Autofit(mode="all"))
        warnings = workbook.report()[0].warnings

    assert [warning.split("]", 1)[0] + "]" for warning in warnings] == [
        "[precision-as-text]",
        "[precision-as-text]",
        "[temporal-as-text]",
        "[temporal-as-text]",
        "[temporal-as-text]",
        "[temporal-as-text]",
    ]
    assert [
        f'column "{name}"' in warning
        for name, warning in zip(
            ("identifier", "decimal", "date", "datetime", "time", "duration"),
            warnings,
            strict=True,
        )
    ] == [True] * 6

    worksheet = openpyxl.load_workbook(output, data_only=False)["Report"]
    assert worksheet["A2"].data_type == "b"
    assert worksheet["B2"].value == 123456789012345
    assert worksheet["B3"].value == "1234567890123456"
    assert worksheet["C2"].value == 1.2
    assert worksheet["C3"].value == "123456789012345678.90"
    assert worksheet["D2"].number_format == "yyyy-mm-dd"
    assert worksheet["D3"].value == "1899-12-31"
    assert worksheet["E2"].number_format == "yyyy-mm-dd hh:mm:ss.000"
    assert worksheet["E3"].value.endswith("123456")
    assert worksheet["F2"].number_format == "hh:mm:ss.000"
    assert worksheet["F3"].value.endswith("123456000")
    assert worksheet["G2"].number_format == "[h]:mm:ss.000"
    assert worksheet["G3"].value == "-1000us"
    assert worksheet["H2"].data_type == "s"
    assert worksheet["H2"].value == "=1+1"


def test_unsupported_dtype_does_not_abort_open_workbook(tmp_path: Path) -> None:
    output = tmp_path / "lifecycle.xlsx"
    unsupported = pl.DataFrame({"payload": pl.Series([[1]], dtype=pl.List(pl.Int64))})

    with nx.Workbook(output) as workbook:
        workbook.write_sheet(pl.DataFrame({"value": [1]}), "Before")
        with pytest.raises(TypeError, match="Unsupported Polars dtype"):
            workbook.write_sheet(unsupported, "Rejected")
        workbook.write_sheet(pl.DataFrame({"value": [2]}), "After")

    assert openpyxl.load_workbook(output).sheetnames == ["Before", "After"]


def test_numeric_cell_kind_is_not_lost_when_role_inference_is_disabled(
    tmp_path: Path,
) -> None:
    output = tmp_path / "no-inference.xlsx"
    with nx.Workbook(output) as workbook:
        workbook.write_sheet(
            pl.DataFrame({"value": [1.25]}),
            "Report",
            infer_numeric_columns=False,
        )

    cell = openpyxl.load_workbook(output, data_only=False)["Report"]["A2"]
    assert cell.data_type == "n"
    assert cell.value == 1.25
