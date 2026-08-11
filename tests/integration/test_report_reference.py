from __future__ import annotations

import json
import math
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import neatxlsx as nx
import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from reference_cases import build_reference_case, scenario_ids

ROOT = Path(__file__).parents[2]
FIXTURE_DIR = ROOT / "tests" / "fixtures" / "report-reference"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _canonical_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        if math.isinf(value):
            return "Inf" if value > 0 else "-Inf"
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    return str(value)


def _cell_manifest(cell: Any) -> dict[str, Any]:
    data_type = cell.data_type
    kind = {
        "b": "boolean",
        "e": "error",
        "f": "formula",
        "n": "blank" if cell.value is None else "number",
    }.get(data_type, "string")
    if cell.value is None:
        kind = "blank"
    return {
        "address": cell.coordinate,
        "kind": kind,
        "value": _canonical_value(cell.value),
        "number_format": cell.number_format,
    }


def _worksheet_manifest(worksheet: Any) -> dict[str, Any]:
    widths: dict[str, float] = {}
    for column_index in range(1, min(worksheet.max_column or 0, 4) + 1):
        letter = get_column_letter(column_index)
        width = worksheet.column_dimensions[letter].width
        if width is not None:
            widths[letter] = round(float(width), 4)
    return {
        "name": worksheet.title,
        "dimensions": [worksheet.max_row or 0, worksheet.max_column or 0],
        "freeze_panes": worksheet.freeze_panes,
        "merges": sorted(str(value) for value in worksheet.merged_cells.ranges),
        "widths": widths,
    }


def _report_manifest(reports: tuple[nx.SheetReport, ...]) -> list[dict[str, Any]]:
    return [
        {
            "requested_name": report.requested_name,
            "worksheets": [
                {
                    "name": part.name,
                    "rows": [part.row_start, part.row_stop],
                    "columns": [part.column_start, part.column_stop],
                }
                for part in report.worksheets
            ],
            "warnings": list(report.warnings),
        }
        for report in reports
    ]


def _assert_ooxml_contract(output: Path, expected: dict[str, Any]) -> None:
    with zipfile.ZipFile(output) as archive:
        names = set(archive.namelist())
        assert {
            "[Content_Types].xml",
            "xl/workbook.xml",
            "xl/styles.xml",
        } <= names
        sheet_names = sorted(
            name
            for name in names
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        assert len(sheet_names) == len(expected["worksheets"])
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        declared_names = [
            element.attrib["name"]
            for element in workbook_root.findall("m:sheets/m:sheet", NS)
        ]
        assert declared_names == [
            worksheet["name"] for worksheet in expected["worksheets"]
        ]
        for sheet_name, worksheet in zip(
            sheet_names, expected["worksheets"], strict=True
        ):
            root = ET.fromstring(archive.read(sheet_name))
            merge_refs = sorted(
                element.attrib["ref"]
                for element in root.findall("m:mergeCells/m:mergeCell", NS)
            )
            assert merge_refs == worksheet["merges"]

    with zipfile.ZipFile(output) as archive:
        for sheet_name, worksheet in zip(
            sheet_names, expected["worksheets"], strict=True
        ):
            formula_like = [
                sentinel
                for sentinel in expected["sentinels"].get(worksheet["name"], [])
                if sentinel["kind"] == "string"
                and isinstance(sentinel["value"], str)
                and sentinel["value"][:1] in {"=", "+"}
            ]
            if not formula_like:
                continue
            root = ET.fromstring(archive.read(sheet_name))
            cells = {
                element.attrib["r"]: element for element in root.findall(".//m:c", NS)
            }
            for sentinel in formula_like:
                cell = cells[sentinel["address"]]
                assert cell.find("m:f", NS) is None
                assert cell.attrib.get("t") in {"inlineStr", "s"}


@pytest.mark.parametrize(
    "scenario_id",
    tuple(
        scenario_id for scenario_id in scenario_ids() if scenario_id != "split-planning"
    ),
)
def test_reference_workbook_matches_manifest(tmp_path: Path, scenario_id: str) -> None:
    case = build_reference_case(scenario_id)
    expected = json.loads(
        (FIXTURE_DIR / f"{scenario_id}.json").read_text(encoding="utf-8")
    )
    output = tmp_path / f"{scenario_id}.xlsx"
    with nx.Workbook(output) as workbook:
        for sheet in case.sheets:
            workbook.write_sheet(sheet.data, sheet.name, **sheet.kwargs)
        reports = workbook.report()

    assert _report_manifest(reports) == expected["reports"]
    assert [
        worksheet.title for worksheet in openpyxl.load_workbook(output).worksheets
    ] == [worksheet["name"] for worksheet in expected["worksheets"]]
    workbook = openpyxl.load_workbook(output, data_only=False)
    assert [
        _worksheet_manifest(worksheet) for worksheet in workbook.worksheets
    ] == expected["worksheets"]
    for sheet_spec in case.sheets:
        worksheet = workbook[sheet_spec.name]
        assert [
            _cell_manifest(worksheet[address]) for address in sheet_spec.sentinels
        ] == expected["sentinels"][sheet_spec.name]
    workbook.close()
    _assert_ooxml_contract(output, expected)
    assert expected["known_gaps"] == list(case.known_gaps)


def test_split_planning_manifest_is_compact_and_columns_first() -> None:
    expected = json.loads(
        (FIXTURE_DIR / "split-planning.json").read_text(encoding="utf-8")
    )
    case = build_reference_case("split-planning")

    assert expected["kind"] == "split-planning"
    assert expected["split_planning"] == case.split_planning
    assert [part["name"] for part in expected["split_planning"]["parts"]] == [
        "Split_1",
        "Split_2",
        "Split_3",
        "Split_4",
    ]
