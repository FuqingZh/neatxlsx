from __future__ import annotations

import os
import stat
import zipfile
from pathlib import Path

import neatxlsx as nx
import neatxlsx.writer as writer_module
import openpyxl
import polars as pl
import pytest


def test_dataframe_write_is_readable_and_reports_physical_sheet(
    tmp_path: Path,
) -> None:
    output = tmp_path / "report.xlsx"
    with nx.Workbook(output) as workbook:
        result = workbook.write_sheet(
            pl.DataFrame({"id": [1, 2], "score": [1.25, 2.5]}),
            "Data",
            integer_columns="id",
            decimal_columns="score",
            freeze_columns=1,
        )
        report = workbook.report()

    assert result is workbook
    assert report == (
        nx.SheetReport(
            requested_name="Data",
            worksheets=(nx.WorksheetPart("Data", 0, 2, 0, 2),),
        ),
    )
    assert workbook.report() == report
    sheet = openpyxl.load_workbook(output, read_only=False)["Data"]
    assert sheet["A2"].value == 1
    assert sheet["B2"].value == 1.25
    assert sheet.freeze_panes == "B2"


@pytest.mark.parametrize("mode", ["none", "header", "body", "all"])
def test_all_autofit_modes_consume_lazyframe_batches_once(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    mode: str,
) -> None:
    output = tmp_path / f"lazy-{mode}.xlsx"
    collect_calls = 0
    sources: list[OneShotBatches] = []
    real_collect_batches = writer_module.collect_batches

    class OneShotBatches:
        def __init__(self, batches: object) -> None:
            self._batches = iter(batches)  # type: ignore[arg-type]
            self.iter_calls = 0
            self.exhausted = False

        def __iter__(self) -> OneShotBatches:
            self.iter_calls += 1
            if self.iter_calls > 1:
                raise AssertionError("batch source was replayed")
            return self

        def __next__(self) -> object:
            try:
                return next(self._batches)
            except StopIteration:
                self.exhausted = True
                raise

    def collect_once(frame: pl.LazyFrame, *, chunk_size: int) -> OneShotBatches:
        nonlocal collect_calls
        collect_calls += 1
        source = OneShotBatches(real_collect_batches(frame, chunk_size=chunk_size))
        sources.append(source)
        return source

    monkeypatch.setattr(writer_module, "collect_batches", collect_once)
    with nx.Workbook(output, chunk_size=1) as workbook:
        workbook.write_sheet(
            pl.LazyFrame(
                {
                    "label": [
                        "short",
                        "a much longer value after the autofit sample",
                        "tail",
                    ]
                }
            ),
            "Data",
            autofit=nx.Autofit(mode=mode, max_rows=1),  # type: ignore[arg-type]
        )

    assert collect_calls == 1
    assert len(sources) == 1
    assert sources[0].iter_calls == 1
    assert sources[0].exhausted is True
    sheet = openpyxl.load_workbook(output)["Data"]
    assert [sheet[f"A{row}"].value for row in range(2, 5)] == [
        "short",
        "a much longer value after the autofit sample",
        "tail",
    ]
    if mode in {"body", "all"}:
        assert sheet.column_dimensions["A"].width < 30


@pytest.mark.parametrize("mode", ["none", "header", "body", "all"])
def test_zero_row_nonempty_schema_succeeds_in_every_autofit_mode(
    tmp_path: Path,
    mode: str,
) -> None:
    output = tmp_path / f"empty-{mode}.xlsx"
    with nx.Workbook(output, chunk_size=1) as workbook:
        workbook.write_sheet(
            pl.DataFrame(schema={"empty": pl.String}).lazy(),
            "Empty",
            autofit=nx.Autofit(mode=mode),  # type: ignore[arg-type]
        )
        report = workbook.report()

    assert report == (
        nx.SheetReport(
            requested_name="Empty",
            worksheets=(nx.WorksheetPart("Empty", 0, 0, 0, 1),),
        ),
    )
    sheet = openpyxl.load_workbook(output)["Empty"]
    assert sheet["A1"].value == "empty"
    assert sheet.max_row == 1


def test_zero_column_input_aborts_workbook(tmp_path: Path) -> None:
    output = tmp_path / "zero-columns.xlsx"
    workbook = nx.Workbook(output)

    with pytest.raises(nx.WriteError, match="zero columns"):
        workbook.write_sheet(pl.DataFrame(), "Zero")

    with pytest.raises(nx.StateError, match="aborted"):
        workbook.write_sheet(pl.DataFrame({"value": [1]}), "After")
    with pytest.raises(nx.StateError, match="aborted"):
        workbook.close()
    assert output.exists() is False


def test_internal_name_collisions_do_not_move_completed_sheets(tmp_path: Path) -> None:
    output = tmp_path / "internal-name-collisions.xlsx"
    requested_names = ["Before", "__nx_tmp_1", "__nx_stage_1", "Before"]
    with nx.Workbook(output) as workbook:
        for value, name in enumerate(requested_names):
            workbook.write_sheet(pl.DataFrame({"value": [value]}), name)
        report_names = [part.worksheets[0].name for part in workbook.report()]

    assert report_names == ["Before", "__nx_tmp_1", "__nx_stage_1", "Before__2"]
    book = openpyxl.load_workbook(output)
    assert book.sheetnames == report_names
    assert [book[name]["A2"].value for name in report_names] == [0, 1, 2, 3]


def test_strings_that_look_like_formulas_remain_literal(tmp_path: Path) -> None:
    output = tmp_path / "literal.xlsx"
    with nx.Workbook(output) as workbook:
        workbook.write_sheet(pl.DataFrame({"text": ["=1+1", "+cmd"]}), "Data")

    sheet = openpyxl.load_workbook(output, data_only=False)["Data"]
    assert sheet["A2"].value == "=1+1"
    assert sheet["A2"].data_type == "s"
    assert sheet["A3"].value == "+cmd"


def test_context_exception_aborts_and_preserves_existing_target(
    tmp_path: Path,
) -> None:
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"original")

    with pytest.raises(RuntimeError, match="stop"):
        with nx.Workbook(output) as workbook:
            workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")
            raise RuntimeError("stop")

    assert output.read_bytes() == b"original"
    assert not list(tmp_path.glob(".*.neatxlsx-*.xlsx"))


def test_close_replaces_target_and_preserves_permission_bits(tmp_path: Path) -> None:
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"old")
    output.chmod(0o640)

    with nx.Workbook(output) as workbook:
        workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")

    assert zipfile.is_zipfile(output)
    assert stat.S_IMODE(output.stat().st_mode) == 0o640


def test_commit_failure_is_retryable_and_keeps_completed_temp(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "report.xlsx"
    workbook = nx.Workbook(output)
    workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")
    real_replace = os.replace
    calls = 0

    def fail_once(
        source: str | bytes | os.PathLike[str] | os.PathLike[bytes],
        target: str | bytes | os.PathLike[str] | os.PathLike[bytes],
    ) -> None:
        nonlocal calls
        calls += 1
        if calls == 1:
            raise PermissionError("denied")
        real_replace(source, target)

    monkeypatch.setattr(writer_module.os, "replace", fail_once)
    with pytest.raises(nx.CommitError) as caught:
        workbook.close()
    assert caught.value.retryable is True
    assert workbook._temp_path.exists()

    workbook.close()
    assert output.exists()


def test_lifecycle_actions_are_idempotent_but_not_interchangeable(
    tmp_path: Path,
) -> None:
    aborted = nx.Workbook(tmp_path / "aborted.xlsx")
    aborted.abort()
    aborted.abort()
    with pytest.raises(nx.StateError):
        aborted.close()

    closed = nx.Workbook(tmp_path / "closed.xlsx")
    closed.write_sheet(pl.DataFrame({"a": [1]}), "Data")
    closed.close()
    closed.close()
    with pytest.raises(nx.StateError):
        closed.abort()
    with pytest.raises(nx.StateError):
        closed.write_sheet(pl.DataFrame({"a": [2]}), "Other")


def test_input_type_error_does_not_poison_open_workbook(tmp_path: Path) -> None:
    output = tmp_path / "recover.xlsx"
    workbook = nx.Workbook(output)
    with pytest.raises(TypeError, match="data must be"):
        workbook.write_sheet({"a": [1]}, "Data")  # type: ignore[arg-type]
    workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")
    workbook.close()

    assert output.exists()


def test_header_value_error_does_not_poison_open_workbook(tmp_path: Path) -> None:
    output = tmp_path / "recover-header.xlsx"
    workbook = nx.Workbook(output)
    with pytest.raises(ValueError, match="header width"):
        workbook.write_sheet(
            pl.DataFrame({"a": [1], "b": [2]}),
            "Data",
            header=pl.DataFrame({"only_one": ["Header"]}),
        )
    workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")
    workbook.close()

    assert output.exists()


def test_zip64_default_and_opt_out_have_distinct_container_versions(
    tmp_path: Path,
) -> None:
    default = tmp_path / "default.xlsx"
    with nx.Workbook(default) as workbook:
        workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")
    classic = tmp_path / "classic.xlsx"
    with nx.Workbook(classic, use_zip64=False) as workbook:
        workbook.write_sheet(pl.DataFrame({"a": [1]}), "Data")

    with zipfile.ZipFile(default) as archive:
        assert min(item.extract_version for item in archive.infolist()) >= 45
    with zipfile.ZipFile(classic) as archive:
        assert max(item.extract_version for item in archive.infolist()) < 45


def test_symbolic_link_target_is_rejected(tmp_path: Path) -> None:
    real = tmp_path / "real.xlsx"
    real.write_bytes(b"content")
    link = tmp_path / "link.xlsx"
    link.symlink_to(real)

    with pytest.raises(ValueError, match="symbolic link"):
        nx.Workbook(link)
