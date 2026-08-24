from __future__ import annotations

import os
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any, cast

import neatxlsx as nx
import openpyxl
import polars as pl
import pytest

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@pytest.mark.parametrize(
    "make_data",
    [
        pytest.param(
            lambda: pl.DataFrame({"identifier": ["ID-1"], "explanation": ["说明"]}),
            id="dataframe",
        ),
        pytest.param(
            lambda: pl.LazyFrame({"identifier": ["ID-1"], "explanation": ["说明"]}),
            id="lazyframe",
        ),
    ],
)
def test_header_rows_and_column_format_patches_reach_xlsx(
    tmp_path: Path,
    make_data: Callable[[], pl.DataFrame | pl.LazyFrame],
) -> None:
    output = tmp_path / "fonts.xlsx"
    header = pl.DataFrame(
        {
            "identifier": ["Metadata", "Identifier"],
            "explanation": ["Readme", "Explanation"],
        }
    )
    with nx.Workbook(
        output,
        text_format=nx.Format(font_name="Times New Roman"),
        header_format=nx.Format(font_name="Times New Roman"),
        use_zip64=False,
    ) as workbook:
        workbook.write_sheet(
            make_data(),
            "Readme",
            header=header,
            header_row_formats=[nx.Format(font_name="SimSun"), None],
            column_formats={"explanation": nx.Format(font_name="SimSun")},
        )

    sheet = openpyxl.load_workbook(output)["Readme"]
    assert sheet["A1"].font.name == "SimSun"
    assert sheet["A2"].font.name == "Times New Roman"
    assert sheet["A3"].font.name == "Times New Roman"
    assert sheet["B3"].font.name == "SimSun"

    with zipfile.ZipFile(output) as archive:
        styles = ET.fromstring(archive.read("xl/styles.xml"))
        worksheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    fonts = styles.findall("m:fonts/m:font", NS)
    cell_formats = styles.findall("m:cellXfs/m:xf", NS)

    def ooxml_font_name(address: str) -> str:
        cell = worksheet.find(f".//m:c[@r='{address}']", NS)
        assert cell is not None
        style_id = int(cell.attrib["s"])
        font_id = int(cell_formats[style_id].attrib["fontId"])
        name = fonts[font_id].find("m:name", NS)
        assert name is not None
        return name.attrib["val"]

    assert ooxml_font_name("A1") == "SimSun"
    assert ooxml_font_name("A2") == "Times New Roman"
    assert ooxml_font_name("A3") == "Times New Roman"
    assert ooxml_font_name("B3") == "SimSun"


def test_invalid_format_rules_fail_before_backend_and_workbook_remains_usable(
    tmp_path: Path,
) -> None:
    output = tmp_path / "recover-format-rules.xlsx"
    data = pl.DataFrame({"identifier": ["ID-1"], "explanation": ["说明"]})
    header = pl.DataFrame(
        {
            "identifier": ["Metadata", "Identifier"],
            "explanation": ["Readme", "Explanation"],
        }
    )
    workbook = nx.Workbook(output)

    with pytest.raises(ValueError, match="length must equal header height"):
        workbook.write_sheet(
            data,
            "BadLength",
            header=header,
            header_row_formats=[nx.Format()],
        )
    with pytest.raises(TypeError, match="items must be"):
        workbook.write_sheet(
            data,
            "BadHeaderType",
            header=header,
            header_row_formats=cast(Any, [nx.Format(), "bad"]),
        )
    with pytest.raises(TypeError, match="must be a mapping"):
        workbook.write_sheet(
            data,
            "BadMapping",
            column_formats=cast(Any, [("explanation", nx.Format())]),
        )
    with pytest.raises(TypeError, match="values must be"):
        workbook.write_sheet(
            data,
            "BadValue",
            column_formats=cast(Any, {"explanation": "bad"}),
        )
    with pytest.raises(ValueError, match="more than once"):
        workbook.write_sheet(
            data,
            "Duplicate",
            column_formats={1: nx.Format(), "explanation": nx.Format()},
        )

    workbook.write_sheet(data, "After")
    assert [report.requested_name for report in workbook.report()] == ["After"]
    workbook.close()
    assert openpyxl.load_workbook(output).sheetnames == ["After"]


@pytest.mark.parametrize(
    ("data", "autofit"),
    [
        pytest.param(
            pl.DataFrame({"score": [0.00000001]}),
            nx.Autofit(mode="header"),
            id="single-pass",
        ),
        pytest.param(
            pl.LazyFrame({"score": [0.00000001]}),
            nx.Autofit(mode="body"),
            id="two-pass",
        ),
    ],
)
def test_column_patch_survives_scientific_overlay_in_all_streaming_paths(
    tmp_path: Path,
    data: pl.DataFrame | pl.LazyFrame,
    autofit: nx.Autofit,
) -> None:
    output = tmp_path / "scientific.xlsx"
    with nx.Workbook(
        output,
        scientific_format=nx.Format(
            font_name="Arial", num_format="0.0E+0", bg_color="#FFF2CC"
        ),
        use_zip64=False,
    ) as workbook:
        workbook.write_sheet(
            data,
            "Data",
            column_formats={"score": nx.Format(font_name="SimSun")},
            scientific_notation=nx.ScientificNotation(scope="decimal"),
            autofit=autofit,
        )

    cell = openpyxl.load_workbook(output)["Data"]["A2"]
    assert cell.font.name == "SimSun"
    assert cell.number_format == "0.0E+0"
    assert cell.fill.fgColor.rgb == "FFFFF2CC"

    with zipfile.ZipFile(output) as archive:
        styles = ET.fromstring(archive.read("xl/styles.xml"))
        worksheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cell_xml = worksheet.find(".//m:c[@r='A2']", NS)
    assert cell_xml is not None
    style_id = int(cell_xml.attrib["s"])
    cell_format = styles.findall("m:cellXfs/m:xf", NS)[style_id]
    font_id = int(cell_format.attrib["fontId"])
    font_name = styles.findall("m:fonts/m:font", NS)[font_id].find("m:name", NS)
    assert font_name is not None
    assert font_name.attrib["val"] == "SimSun"
    num_format_id = int(cell_format.attrib["numFmtId"])
    num_formats = {
        int(item.attrib["numFmtId"]): item.attrib["formatCode"]
        for item in styles.findall("m:numFmts/m:numFmt", NS)
    }
    assert num_formats[num_format_id] == "0.0E+0"


@pytest.mark.parametrize(
    ("data", "autofit"),
    [
        pytest.param(
            pl.DataFrame({"score": [0.00000001]}),
            nx.Autofit(mode="header"),
            id="single-pass",
        ),
        pytest.param(
            pl.LazyFrame({"score": [0.00000001]}),
            nx.Autofit(mode="body"),
            id="two-pass",
        ),
    ],
)
def test_omitted_column_formats_keep_the_original_scientific_style(
    tmp_path: Path,
    data: pl.DataFrame | pl.LazyFrame,
    autofit: nx.Autofit,
) -> None:
    output = tmp_path / "scientific-compatibility.xlsx"
    with nx.Workbook(
        output,
        decimal_format=nx.Format(
            font_name="Courier New", bold=True, bg_color="#FF0000"
        ),
        scientific_format=nx.Format(font_name="Arial", num_format="0.0E+0"),
        use_zip64=False,
    ) as workbook:
        workbook.write_sheet(
            data,
            "Data",
            scientific_notation=nx.ScientificNotation(scope="decimal"),
            autofit=autofit,
        )

    cell = openpyxl.load_workbook(output)["Data"]["A2"]
    assert cell.font.name == "Arial"
    assert cell.font.bold is False
    assert cell.number_format == "0.0E+0"
    assert cell.fill.patternType is None

    with zipfile.ZipFile(output) as archive:
        styles = ET.fromstring(archive.read("xl/styles.xml"))
        worksheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cell_xml = worksheet.find(".//m:c[@r='A2']", NS)
    assert cell_xml is not None
    style_id = int(cell_xml.attrib["s"])
    cell_format = styles.findall("m:cellXfs/m:xf", NS)[style_id]
    font_id = int(cell_format.attrib["fontId"])
    font = styles.findall("m:fonts/m:font", NS)[font_id]
    font_name = font.find("m:name", NS)
    assert font_name is not None
    assert font_name.attrib["val"] == "Arial"
    assert font.find("m:b", NS) is None


@pytest.mark.skipif(
    os.environ.get("NEATXLSX_INCLUDE_LARGE") != "1",
    reason="real Excel-limit pagination runs only in the explicit release gate",
)
def test_format_overrides_survive_physical_row_pagination(tmp_path: Path) -> None:
    output = tmp_path / "split-formats.xlsx"
    data = pl.select(pl.int_range(0, 1_048_575, eager=False).alias("identifier")).lazy()
    header = pl.DataFrame({"identifier": ["Metadata", "Identifier"]})
    with nx.Workbook(
        output,
        integer_format=nx.Format(font_name="Times New Roman"),
        header_format=nx.Format(font_name="Times New Roman"),
    ) as workbook:
        workbook.write_sheet(
            data,
            "Data",
            header=header,
            header_row_formats=[nx.Format(font_name="SimSun"), None],
            column_formats={"identifier": nx.Format(font_name="SimSun")},
            autofit=nx.Autofit(mode="none"),
        )
        parts = workbook.report()[0].worksheets

    assert len(parts) == 2
    book = openpyxl.load_workbook(output, read_only=True)
    for part in parts:
        sheet = book[part.name]
        assert sheet["A1"].font.name == "SimSun"
        assert sheet["A2"].font.name == "Times New Roman"
        assert sheet["A3"].font.name == "SimSun"
