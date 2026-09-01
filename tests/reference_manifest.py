"""Canonical semantic and OOXML projections for XLSX reference fixtures."""

from __future__ import annotations

import math
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import asdict, is_dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import neatxlsx as nx
import openpyxl
import polars as pl
from openpyxl.utils import get_column_letter

from tests.reference_cases import MANIFEST_SCHEMA_VERSION, ReferenceCase

BASELINE_PROVENANCE = {
    "bridge_abi": 5,
    "bridge_contract": "neatxlsx.xlsx.writer.v5",
    "package_version": "0.2.0",
    "source_revision": "85c3519e7fd830c83dd0b3e52b00ff8db19c4ab3",
}

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS = {"m": _MAIN_NS, "r": _DOCUMENT_REL_NS, "pr": _PACKAGE_REL_NS}


def _canonical_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        if math.isinf(value):
            return "Inf" if value > 0 else "-Inf"
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    if is_dataclass(value) and not isinstance(value, type):
        return {key: _canonical_value(item) for key, item in asdict(value).items()}
    if isinstance(value, pl.DataFrame):
        return {
            "kind": "DataFrame",
            "schema": [[name, str(dtype)] for name, dtype in value.schema.items()],
            "rows": [
                {key: _canonical_value(item) for key, item in row.items()}
                for row in value.to_dicts()
            ],
        }
    if isinstance(value, dict):
        return {
            str(key): _canonical_value(item)
            for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))
        }
    if isinstance(value, (list, tuple)):
        return [_canonical_value(item) for item in value]
    return str(value)


def _case_input_manifest(case: ReferenceCase) -> dict[str, Any]:
    return {
        "workbook_kwargs": _canonical_value(case.workbook_kwargs or {}),
        "sheets": [
            {
                "data_kind": type(sheet.data).__name__,
                "name": sheet.name,
                "schema": [
                    [name, str(dtype)]
                    for name, dtype in sheet.data.collect_schema().items()
                ],
                "kwargs": _canonical_value(sheet.kwargs),
                "sentinels": list(sheet.sentinels),
            }
            for sheet in case.sheets
        ],
    }


def _cell_manifest(cell: Any) -> dict[str, Any]:
    data_type = cell.data_type
    kind = {
        "b": "boolean",
        "e": "error",
        "f": "formula",
        "n": "blank" if cell.value is None else "number",
    }.get(data_type, "string")
    if cell.value is None:
        kind = "blank"
    return {
        "address": cell.coordinate,
        "kind": kind,
        "value": _canonical_value(cell.value),
        "number_format": cell.number_format,
    }


def _worksheet_manifest(
    worksheet: Any, *, include_widths: bool = True
) -> dict[str, Any]:
    max_column = worksheet.max_column or 0
    widths: dict[str, float] = {}
    if include_widths:
        for column_index in range(1, max_column + 1):
            letter = get_column_letter(column_index)
            width = worksheet.column_dimensions[letter].width
            if width is not None:
                widths[letter] = round(float(width), 4)
    merged_cells = getattr(worksheet, "merged_cells", None)
    return {
        "name": worksheet.title,
        "dimensions": [worksheet.max_row or 0, max_column],
        "freeze_panes": getattr(worksheet, "freeze_panes", None),
        "merges": sorted(str(value) for value in merged_cells.ranges)
        if merged_cells is not None
        else [],
        "widths": widths,
    }


def _report_manifest(reports: tuple[nx.SheetReport, ...]) -> list[dict[str, Any]]:
    return [
        {
            "requested_name": report.requested_name,
            "worksheets": [
                {
                    "name": part.name,
                    "rows": [part.row_start, part.row_stop],
                    "columns": [part.column_start, part.column_stop],
                }
                for part in report.worksheets
            ],
            "warnings": list(report.warnings),
        }
        for report in reports
    ]


def _relationship_target(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return f"xl/{target}"


def _selected_attributes(element: ET.Element, names: tuple[str, ...]) -> dict[str, str]:
    return {name: element.attrib[name] for name in names if name in element.attrib}


def _ooxml_manifest(
    output: Path,
    sentinels: dict[str, list[dict[str, Any]]],
) -> dict[str, Any]:
    with zipfile.ZipFile(output) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relationship.attrib["Id"]: _relationship_target(
                relationship.attrib["Target"]
            )
            for relationship in relationships_root.findall("pr:Relationship", _NS)
        }
        workbook_sheets = []
        worksheet_projections = []
        for sheet in workbook_root.findall("m:sheets/m:sheet", _NS):
            relationship_id = sheet.attrib[f"{{{_DOCUMENT_REL_NS}}}id"]
            target = targets[relationship_id]
            workbook_sheets.append(
                {
                    "name": sheet.attrib["name"],
                    "relationship_id": relationship_id,
                    "target": target,
                }
            )
            root = ET.fromstring(archive.read(target))
            dimension = root.find("m:dimension", _NS)
            columns = [
                _selected_attributes(
                    element,
                    (
                        "min",
                        "max",
                        "width",
                        "style",
                        "hidden",
                        "bestFit",
                        "customWidth",
                    ),
                )
                for element in root.findall("m:cols/m:col", _NS)
            ]
            merge_refs = sorted(
                element.attrib["ref"]
                for element in root.findall("m:mergeCells/m:mergeCell", _NS)
            )
            formula_addresses = {
                item["address"]
                for item in sentinels.get(sheet.attrib["name"], [])
                if item["kind"] == "string"
                and isinstance(item["value"], str)
                and item["value"][:1] in {"=", "+"}
            }
            cells = {
                element.attrib["r"]: element for element in root.findall(".//m:c", _NS)
            }
            formula_like = [
                {
                    "address": address,
                    "has_formula": cells[address].find("m:f", _NS) is not None,
                    "type": cells[address].attrib.get("t"),
                }
                for address in sorted(formula_addresses)
            ]
            worksheet_projections.append(
                {
                    "columns": columns,
                    "dimension_ref": dimension.attrib.get("ref")
                    if dimension is not None
                    else None,
                    "formula_like_cells": formula_like,
                    "merges": merge_refs,
                    "name": sheet.attrib["name"],
                    "target": target,
                }
            )

        styles_root = ET.fromstring(archive.read("xl/styles.xml"))
        num_formats = [
            _selected_attributes(element, ("numFmtId", "formatCode"))
            for element in styles_root.findall("m:numFmts/m:numFmt", _NS)
        ]
        cell_xfs = [
            dict(sorted(element.attrib.items()))
            for element in styles_root.findall("m:cellXfs/m:xf", _NS)
        ]

    return {
        "styles": {"cell_xfs": cell_xfs, "number_formats": num_formats},
        "workbook_sheets": workbook_sheets,
        "worksheets": worksheet_projections,
    }


def manifest_from_workbook(
    output: Path,
    case: ReferenceCase,
    reports: tuple[nx.SheetReport, ...],
) -> dict[str, Any]:
    """Build a deterministic compatibility manifest for one scenario."""
    common = {
        "baseline": dict(BASELINE_PROVENANCE),
        "inputs": _case_input_manifest(case),
        "known_gaps": [dict(gap) for gap in case.known_gaps],
        "scenario": case.scenario_id,
        "schema_version": MANIFEST_SCHEMA_VERSION,
    }
    if case.split_planning is not None:
        return {
            **common,
            "kind": "split-planning",
            "split_planning": case.split_planning,
        }

    read_only = case.scenario_id == "large-row-split"
    workbook = openpyxl.load_workbook(
        output,
        data_only=False,
        read_only=read_only,
    )
    sentinels: dict[str, list[dict[str, Any]]] = {}
    for sheet_spec in case.sheets:
        if not sheet_spec.sentinels:
            continue
        worksheet = workbook[sheet_spec.name]
        sentinels[sheet_spec.name] = [
            _cell_manifest(worksheet[address]) for address in sheet_spec.sentinels
        ]
    worksheets = [
        _worksheet_manifest(worksheet, include_widths=not read_only)
        for worksheet in workbook.worksheets
    ]
    workbook.close()

    manifest = {
        **common,
        "kind": "workbook",
        "reports": _report_manifest(reports),
        "sentinels": sentinels,
        "worksheets": worksheets,
    }
    if not read_only:
        manifest["ooxml"] = _ooxml_manifest(output, sentinels)
    return manifest
