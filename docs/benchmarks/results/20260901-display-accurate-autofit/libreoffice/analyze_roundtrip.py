from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "display-autofit-acceptance-nonzip64.xlsx"
ROUNDTRIP = ROOT / "roundtrip-nonzip64" / "display-autofit-acceptance-nonzip64.xlsx"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cols(path: Path) -> dict[str, list[dict[str, str]]]:
    result: dict[str, list[dict[str, str]]] = {}
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            target = targets[rid]
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            root = ET.fromstring(archive.read(target))
            result[sheet.attrib["name"]] = [
                {k: v for k, v in col.attrib.items() if k in {"min", "max", "width", "bestFit", "customWidth", "hidden", "style"}}
                for col in root.findall("m:cols/m:col", NS)
            ]
    return result


def workbook_projection(path: Path) -> tuple[dict[str, dict[str, tuple[object, str, str]]], dict[str, list[str]]]:
    workbook = load_workbook(path, data_only=False)
    cells: dict[str, dict[str, tuple[object, str, str]]] = {}
    merges: dict[str, list[str]] = {}
    for sheet in workbook.worksheets:
        cells[sheet.title] = {
            cell.coordinate: (cell.value, cell.data_type, cell.number_format)
            for row in sheet.iter_rows() for cell in row if cell.value is not None
        }
        merges[sheet.title] = sorted(str(item) for item in sheet.merged_cells.ranges)
    return cells, merges


source_cells, source_merges = workbook_projection(SOURCE)
roundtrip_cells, roundtrip_merges = workbook_projection(ROUNDTRIP)
source_cols, roundtrip_cols = cols(SOURCE), cols(ROUNDTRIP)
value_diffs = {}
format_diffs = {}
type_diffs = {}
for sheet in source_cells:
    for address in sorted(set(source_cells[sheet]) | set(roundtrip_cells.get(sheet, {}))):
        before, after = source_cells[sheet].get(address), roundtrip_cells.get(sheet, {}).get(address)
        if before is None or after is None or before[0] != after[0]:
            value_diffs.setdefault(sheet, []).append(address)
        if before is None or after is None or before[1] != after[1]:
            type_diffs.setdefault(sheet, []).append(address)
        if before is None or after is None or before[2] != after[2]:
            format_diffs.setdefault(sheet, []).append(address)
report = {
    "input": {"path": SOURCE.name, "sha256": digest(SOURCE)},
    "roundtrip": {"path": str(ROUNDTRIP.relative_to(ROOT)), "sha256": digest(ROUNDTRIP)},
    "sheets_equal": list(source_cells) == list(roundtrip_cells),
    "value_diffs": value_diffs,
    "type_diffs": type_diffs,
    "format_diffs": format_diffs,
    "merges_equal": source_merges == roundtrip_merges,
    "source_merges": source_merges,
    "roundtrip_merges": roundtrip_merges,
    "source_cols": source_cols,
    "roundtrip_cols": roundtrip_cols,
    "source_bestfit_count": sum(col.get("bestFit") == "1" for group in source_cols.values() for col in group),
    "roundtrip_bestfit_count": sum(col.get("bestFit") == "1" for group in roundtrip_cols.values() for col in group),
}
(ROOT / "roundtrip-analysis.json").write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
print(json.dumps({k: report[k] for k in ("sheets_equal", "merges_equal", "value_diffs", "type_diffs", "format_diffs", "source_bestfit_count", "roundtrip_bestfit_count")}, ensure_ascii=False, indent=2))
