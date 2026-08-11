#!/usr/bin/env python3
"""Generate and verify deterministic M0 report-reference artifacts."""

from __future__ import annotations

import argparse
import difflib
import json
import math
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import neatxlsx as nx
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = ROOT / "tests" / "fixtures" / "report-reference"

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from tests.reference_cases import (  # noqa: E402
    MANIFEST_SCHEMA_VERSION,
    ReferenceCase,
    build_reference_case,
    scenario_ids,
)


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate and verify neatxlsx report-reference artifacts."
    )
    parser.add_argument(
        "--scenario",
        choices=scenario_ids(include_large=True),
        action="append",
        dest="scenarios",
        help="Generate one scenario; repeat to select several.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("dist") / "report-reference",
    )
    parser.add_argument(
        "--include-large",
        action="store_true",
        help="Include the optional real row-limit XLSX artifact.",
    )
    parser.add_argument(
        "--accept",
        action="store_true",
        help="Replace expected JSON fixtures with the generated manifests.",
    )
    return parser.parse_args()


def _canonical_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return "NaN"
        if math.isinf(value):
            return "Inf" if value > 0 else "-Inf"
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return value.total_seconds()
    return str(value)


def _cell_manifest(cell: Any) -> dict[str, Any]:
    data_type = cell.data_type
    kind = {
        "b": "boolean",
        "e": "error",
        "f": "formula",
        "n": "blank" if cell.value is None else "number",
    }.get(data_type, "string")
    if cell.value is None:
        kind = "blank"
    return {
        "address": cell.coordinate,
        "kind": kind,
        "value": _canonical_value(cell.value),
        "number_format": cell.number_format,
    }


def _worksheet_manifest(
    worksheet: Any, *, include_widths: bool = True
) -> dict[str, Any]:
    max_column = worksheet.max_column or 0
    widths: dict[str, float] = {}
    if include_widths:
        for column_index in range(1, min(max_column, 4) + 1):
            letter = get_column_letter(column_index)
            width = worksheet.column_dimensions[letter].width
            if width is not None:
                widths[letter] = round(float(width), 4)
    merged_cells = getattr(worksheet, "merged_cells", None)
    return {
        "name": worksheet.title,
        "dimensions": [worksheet.max_row or 0, max_column],
        "freeze_panes": getattr(worksheet, "freeze_panes", None),
        "merges": sorted(str(value) for value in merged_cells.ranges)
        if merged_cells is not None
        else [],
        "widths": widths,
    }


def _report_manifest(reports: tuple[nx.SheetReport, ...]) -> list[dict[str, Any]]:
    return [
        {
            "requested_name": report.requested_name,
            "worksheets": [
                {
                    "name": part.name,
                    "rows": [part.row_start, part.row_stop],
                    "columns": [part.column_start, part.column_stop],
                }
                for part in report.worksheets
            ],
            "warnings": list(report.warnings),
        }
        for report in reports
    ]


def manifest_from_workbook(
    output: Path,
    case: ReferenceCase,
    reports: tuple[nx.SheetReport, ...],
) -> dict[str, Any]:
    """Build the compact manifest for one generated workbook."""
    if case.split_planning is not None:
        return {
            "schema_version": MANIFEST_SCHEMA_VERSION,
            "scenario": case.scenario_id,
            "kind": "split-planning",
            "split_planning": case.split_planning,
            "known_gaps": [dict(gap) for gap in case.known_gaps],
        }

    read_only = case.scenario_id == "large-row-split"
    workbook = openpyxl.load_workbook(
        output,
        data_only=False,
        read_only=read_only,
    )
    sentinels: dict[str, list[dict[str, Any]]] = {}
    for sheet_spec in case.sheets:
        if not sheet_spec.sentinels:
            continue
        worksheet = workbook[sheet_spec.name]
        sentinels[sheet_spec.name] = [
            _cell_manifest(worksheet[address]) for address in sheet_spec.sentinels
        ]
    manifest = {
        "schema_version": MANIFEST_SCHEMA_VERSION,
        "scenario": case.scenario_id,
        "kind": "workbook",
        "reports": _report_manifest(reports),
        "worksheets": [
            _worksheet_manifest(worksheet, include_widths=not read_only)
            for worksheet in workbook.worksheets
        ],
        "sentinels": sentinels,
        "known_gaps": [dict(gap) for gap in case.known_gaps],
    }
    workbook.close()
    return manifest


def _write_case_workbook(
    output: Path, case: ReferenceCase
) -> tuple[nx.SheetReport, ...]:
    output.parent.mkdir(parents=True, exist_ok=True)
    with nx.Workbook(output) as workbook:
        for sheet in case.sheets:
            workbook.write_sheet(sheet.data, sheet.name, **sheet.kwargs)
        return workbook.report()


def generate_scenario(
    scenario_id: str,
    *,
    output_dir: Path,
    fixture_dir: Path = FIXTURE_DIR,
    accept: bool = False,
) -> dict[str, Any]:
    """Generate one scenario and compare or explicitly accept its manifest."""
    case = build_reference_case(scenario_id)
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{scenario_id}.xlsx"
    reports = _write_case_workbook(output, case) if case.sheets else ()
    manifest = manifest_from_workbook(output, case, reports)
    actual_path = output_dir / f"{scenario_id}.json"
    actual_path.write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    expected_path = fixture_dir / f"{scenario_id}.json"
    if accept:
        expected_path.parent.mkdir(parents=True, exist_ok=True)
        expected_path.write_text(
            actual_path.read_text(encoding="utf-8"), encoding="utf-8"
        )
    elif not expected_path.exists():
        raise RuntimeError(
            f"Missing expected fixture {expected_path}; rerun with --accept."
        )
    else:
        expected = expected_path.read_text(encoding="utf-8").splitlines(keepends=True)
        actual = actual_path.read_text(encoding="utf-8").splitlines(keepends=True)
        if expected != actual:
            diff = "".join(
                difflib.unified_diff(
                    expected,
                    actual,
                    fromfile=str(expected_path),
                    tofile=str(actual_path),
                )
            )
            raise RuntimeError(f"Reference manifest mismatch:\n{diff}")
    artifacts = [str(actual_path)]
    if output.exists():
        artifacts.insert(0, str(output))
    print(f"{scenario_id}: {', '.join(artifacts)}")
    return manifest


def main() -> None:
    args = _parse_args()
    selected = tuple(args.scenarios or scenario_ids(include_large=args.include_large))
    if "large-row-split" in selected and not args.include_large:
        raise SystemExit("large-row-split requires --include-large")
    output_dir = args.output_dir
    for scenario_id in selected:
        generate_scenario(
            scenario_id,
            output_dir=output_dir,
            accept=args.accept,
        )


if __name__ == "__main__":
    main()
